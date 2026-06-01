import pytest
from pathlib import Path
from backend.services.document_parser import parse_document, ParsedDocument

FIXTURES_DIR = Path(__file__).parent / "fixtures"

def test_parse_excel_returns_text_and_tables():
    import openpyxl
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["항목명", "규격", "수량", "단위", "단가", "합계"])
    ws.append(["타일공사", "300x300", 100, "장", 5000, 500000])
    ws.append(["도배공사", "합지", 50, "롤", 30000, 1500000])
    wb.save(FIXTURES_DIR / "sample.xlsx")

    result = parse_document(str(FIXTURES_DIR / "sample.xlsx"))

    assert isinstance(result, ParsedDocument)
    assert "타일공사" in result.text
    assert len(result.tables) > 0
    assert result.tables[0][1][0] == "타일공사"

def test_parse_unsupported_format_raises():
    with pytest.raises(ValueError, match="Unsupported"):
        parse_document("file.docx")

def test_parsed_document_has_file_type():
    import openpyxl
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["test", "data"])
    wb.save(FIXTURES_DIR / "type_test.xlsx")

    result = parse_document(str(FIXTURES_DIR / "type_test.xlsx"))
    assert result.file_type == "excel"
